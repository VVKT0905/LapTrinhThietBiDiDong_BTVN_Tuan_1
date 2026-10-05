import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import os

doc = Document()

# Set standard margins (1 inch)
for section in doc.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

def set_para_font(p, name='Times New Roman', size=13, bold=False, italic=False, color=None, align=WD_ALIGN_PARAGRAPH.LEFT):
    p.alignment = align
    for r in p.runs:
        r.font.name = name
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        if color:
            r.font.color.rgb = color

# Header / Institution
p_inst = doc.add_paragraph('TRƯỜNG ĐẠI HỌC GIAO THÔNG VẬN TẢI TP. HỒ CHÍ MINH\nBAN CÔNG NGHỆ SỐ - KHOA CÔNG NGHỆ THÔNG TIN')
set_para_font(p_inst, name='Times New Roman', size=12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

p_div = doc.add_paragraph('-----------------o0o-----------------')
set_para_font(p_div, name='Times New Roman', size=11, align=WD_ALIGN_PARAGRAPH.CENTER)

# Document Title
p_title = doc.add_paragraph('\nBÁO CÁO BÀI TẬP VỀ NHÀ TUẦN 1\nMÔN: LẬP TRÌNH THIẾT BỊ DI ĐỘNG')
set_para_font(p_title, name='Times New Roman', size=16, bold=True, color=RGBColor(0, 51, 102), align=WD_ALIGN_PARAGRAPH.CENTER)

# Metadata Table
table = doc.add_table(rows=6, cols=2)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False

meta_data = [
    ('Giảng viên hướng dẫn:', 'ThS. Trương Quang Tuấn'),
    ('Mã lớp học phần:', '[012012103402] - Lập trình thiết bị di động - CNS_CS1'),
    ('Sinh viên thực hiện:', 'Võ Văn Khương Thịnh'),
    ('Mã số sinh viên (MSSV):', '049206001215'),
    ('Email sinh viên:', 'thinhvvk1215@ut.edu.vn'),
    ('GitHub Repository:', 'https://github.com/VVKT0905/LapTrinhThietBiDiDong_BTVN_Tuan_1')
]

for idx, (lbl, val) in enumerate(meta_data):
    r = table.rows[idx]
    c1, c2 = r.cells[0], r.cells[1]
    c1.width = Inches(2.2)
    c2.width = Inches(4.3)
    p1 = c1.paragraphs[0]
    p1.text = lbl
    set_para_font(p1, name='Times New Roman', size=11.5, bold=True)
    p2 = c2.paragraphs[0]
    p2.text = val
    set_para_font(p2, name='Times New Roman', size=11.5, bold=False, color=RGBColor(0, 70, 150) if 'http' in val else None)

doc.add_paragraph('\n')

# 1. MÔ TẢ NGẮN GỌN VỀ BÀI TẬP
h1 = doc.add_heading('1. Mô tả ngắn gọn về bài tập', level=1)
set_para_font(h1, name='Times New Roman', size=13.5, bold=True, color=RGBColor(0, 51, 102))

p1 = doc.add_paragraph(
    'Bài tập tuần 1 gồm hai nội dung chính:\n'
    '1. Thực hành (Slide 52): Viết app Flutter hiển thị màn hình Profile cá nhân theo đúng giao diện mẫu (nút Back, nút Edit, Avatar tròn ở giữa, Họ tên, MSSV) và đẩy toàn bộ code lên GitHub.\n'
    '2. Lý thuyết (Slide 52 & 53): Trình bày định hướng học tập cá nhân, nhận định sự phát triển của lập trình di động, tìm hiểu mô hình giáo dục HAA và cách học trong kỷ nguyên AI.'
)
set_para_font(p1, name='Times New Roman', size=12)

# 2. MỤC TIÊU ĐẠT ĐƯỢC
h2 = doc.add_heading('2. Mục tiêu đạt được', level=1)
set_para_font(h2, name='Times New Roman', size=13.5, bold=True, color=RGBColor(0, 51, 102))

p2 = doc.add_paragraph(
    '• Nắm được quy trình tạo project, chia widget và quản lý trạng thái cơ bản trên Flutter/Dart.\n'
    '• Tự dựng giao diện theo mẫu, xử lý ảnh đại diện tròn, tạo dialog chỉnh sửa thông tin và bottom sheet hiển thị thông tin học phần.\n'
    '• Đóng gói và nộp bài đúng cấu trúc thư mục quy định.'
)
set_para_font(p2, name='Times New Roman', size=12)

# 3. KẾT QUẢ ĐẠT ĐƯỢC & HÌNH ẢNH ĐẦU RA (OUTPUT)
h3 = doc.add_heading('3. Kết quả đạt được & Hình ảnh đầu ra (Output)', level=1)
set_para_font(h3, name='Times New Roman', size=13.5, bold=True, color=RGBColor(0, 51, 102))

p3 = doc.add_paragraph(
    'Ứng dụng đã hoàn thành đúng theo mẫu, chạy ổn định trên Windows và Web.'
)
set_para_font(p3, name='Times New Roman', size=12)

img_path = os.path.join(os.path.dirname(__file__), '..', '..', 'TaiLieu', 'screenshots', 'app_preview.png')
if not os.path.exists(img_path):
    img_path = 'TaiLieu/screenshots/app_preview.png'

if os.path.exists(img_path):
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_picture(img_path, width=Inches(2.5))
    p_cap = doc.add_paragraph('Hình 1: Giao diện ứng dụng Profile Sinh Viên chạy thực tế')
    set_para_font(p_cap, name='Times New Roman', size=10.5, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER)

# 4. GIẢI THÍCH CÁC HÀM VÀ WIDGET CHÍNH
h4 = doc.add_heading('4. Giải thích các hàm và widget chính trong source code', level=1)
set_para_font(h4, name='Times New Roman', size=13.5, bold=True, color=RGBColor(0, 51, 102))

p4 = doc.add_paragraph(
    'Mã nguồn chính nằm tại file SourceCode/lib/main.dart:\n'
    '• StudentProfileApp: Khởi tạo app và cấu hình theme Material 3.\n'
    '• ProfileScreen (_ProfileScreenState): Quản lý state thông tin hiển thị của sinh viên (Họ tên, MSSV).\n'
    '• _showEditDialog(): Mở AlertDialog cho phép nhập và cập nhật ngay Họ tên/MSSV mới qua setState().\n'
    '• _showExtraInfoSheet(): Mở ModalBottomSheet hiển thị thông tin học phần mở rộng (Khoa, Lớp tín chỉ, Email sinh viên) khi chạm vào avatar hoặc nút chi tiết.\n'
    '• ClipOval & Image.asset: Cắt ảnh chân dung thành hình tròn, có fallback hiển thị icon nếu không load được ảnh.'
)
set_para_font(p4, name='Times New Roman', size=12)

# 5. TRẢ LỜI CÂU HỎI LÝ THUYẾT (SLIDE 52 & SLIDE 53)
h5 = doc.add_heading('5. Nội dung trả lời câu hỏi lý thuyết (Slide 52 & 53)', level=1)
set_para_font(h5, name='Times New Roman', size=13.5, bold=True, color=RGBColor(0, 51, 102))

# 5.1
h5_1 = doc.add_heading('5.1. Định hướng cá nhân sau môn học (Slide 52 - Câu 1)', level=2)
set_para_font(h5_1, name='Times New Roman', size=12, bold=True)
p5_1 = doc.add_paragraph(
    'Em định hướng theo con đường Fullstack Developer. Em muốn học chắc mảng Mobile (Flutter) để có thể tự tay làm trọn vẹn cả phần giao diện ứng dụng (Client) lẫn phần máy chủ (Backend), không bị giới hạn vào một nền tảng duy nhất.\n'
    'Mục tiêu của em trong môn học là nắm chắc kiến trúc app, cách gọi API và quản lý state để làm đồ án tốt nghiệp và có sản phẩm đưa vào CV ứng tuyển thực tập.'
)
set_para_font(p5_1, name='Times New Roman', size=12)

# 5.2
h5_2 = doc.add_heading('5.2. Trong 10 năm tới lập trình di động có phát triển không? (Slide 52 - Câu 2)', level=2)
set_para_font(h5_2, name='Times New Roman', size=12, bold=True)
p5_2 = doc.add_paragraph(
    'Em khẳng định là vẫn phát triển rất mạnh, vì các lý do sau:\n'
    '1. Smartphone vẫn là vật bất ly thân: Mọi nhu cầu của con người từ thanh toán, định danh cá nhân (VNeID), ví giấy tờ, mạng xã hội và giải trí đều nằm trên điện thoại.\n'
    '2. Làn sóng On-device AI (Edge AI): Chip điện thoại thế hệ mới chạy được AI trực tiếp trên máy mà không cần gửi dữ liệu lên server, vừa nhanh vừa bảo mật. Lập trình viên di động sẽ chuyển từ làm giao diện tĩnh sang làm app thông minh tương tác theo ngữ cảnh.\n'
    '3. Điều khiển hệ sinh thái IoT: Điện thoại là trung tâm điều khiển cho nhà thông minh, xe ô tô và các thiết bị đeo cá nhân.'
)
set_para_font(p5_2, name='Times New Roman', size=12)

# 5.3
h5_3 = doc.add_heading('5.3. Mô hình giáo dục HAA (Horowitz Andreessen Academy) (Slide 53 - Câu 1)', level=2)
set_para_font(h5_3, name='Times New Roman', size=12, bold=True)
p5_3 = doc.add_paragraph(
    '• Bản chất: Bắt nguồn từ triết lý của quỹ đầu tư a16z: "Internet → Infinite knowledge. AI → Infinite doing. Education must evolve from knowing to doing." Khi Internet mang lại tri thức vô tận và AI mang lại năng lực thực thi vô tận, giáo dục không thể dừng lại ở việc học thuộc kiến thức (knowing) mà phải chuyển sang làm sản phẩm thật (doing).\n'
    '• Ưu điểm: Học qua thực chiến (Learn by building), bám sát thị trường, ứng dụng AI để tăng tốc độ làm việc.\n'
    '• Nhược điểm: Dễ bị hổng kiến thức khoa học máy tính nền tảng (cấu trúc dữ liệu, giải thuật, hệ điều hành) nếu quá lạm dụng code AI sinh sẵn; dễ ngộ nhận năng lực thực tế của bản thân.'
)
set_para_font(p5_3, name='Times New Roman', size=12)

# 5.4 Table
h5_4 = doc.add_heading('5.4. So sánh cách tiếp cận tri thức qua 3 thời kỳ (Slide 53 - Câu 2)', level=2)
set_para_font(h5_4, name='Times New Roman', size=12, bold=True)

t_compare = doc.add_table(rows=4, cols=4)
t_compare.alignment = WD_TABLE_ALIGNMENT.CENTER

headers = ['Tiêu chí', '1. Trước Internet', '2. Thời kỳ Internet', '3. Kỷ nguyên AI hiện nay']
for j, h in enumerate(headers):
    cell = t_compare.rows[0].cells[j]
    cell.paragraphs[0].text = h
    set_para_font(cell.paragraphs[0], name='Times New Roman', size=11, bold=True, color=RGBColor(255, 255, 255), align=WD_ALIGN_PARAGRAPH.CENTER)
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="003366"/>')
    cell._tc.get_or_add_tcPr().append(shd)

rows_data = [
    ('Tìm kiếm',
     'Tốn công: Tìm sách thư viện, hỏi thầy cô.',
     'Dễ dàng: Tìm bằng từ khóa trên Google, Stack Overflow.',
     'Trực tiếp: Hỏi đáp tự nhiên, AI trả lời trúng đích theo ngữ cảnh.'),
    ('Ghi nhớ',
     'Não bộ phải nhớ từng cú pháp, công thức.',
     'Nhớ từ khóa và nơi lưu tài liệu.',
     'Não bộ giải phóng việc nhớ cú pháp, chỉ cần nhớ nguyên lý cốt lõi.'),
    ('Vận dụng',
     'Làm thủ công từng dòng lệnh, sửa lỗi rất lâu.',
     'Copy-paste code mẫu rồi tự sửa lại.',
     'Ra lệnh cho AI tạo khung, người lập trình kiểm tra, sửa lỗi và tối ưu.')
]

for i, row in enumerate(rows_data):
    for j, val in enumerate(row):
        cell = t_compare.rows[i+1].cells[j]
        cell.paragraphs[0].text = val
        bold_flag = (j == 0)
        set_para_font(cell.paragraphs[0], name='Times New Roman', size=10.5, bold=bold_flag)

doc.add_paragraph('\n')

# 5.5
h5_5 = doc.add_heading('5.5. Cần học gì khi AI có thể làm gần như mọi thứ? (Slide 53 - Câu 3)', level=2)
set_para_font(h5_5, name='Times New Roman', size=12, bold=True)
p5_5 = doc.add_paragraph(
    '1. Kiến trúc hệ thống (System Design): AI viết từng hàm rất nhanh, nhưng người kỹ sư phải quyết định ghép nối các phần sao cho an toàn, dễ bảo trì và mở rộng.\n'
    '2. Kỹ năng kiểm tra và phản biện code (Verification): Biết đọc hiểu, phát hiện bug ngầm và lỗ hổng bảo mật trong code do AI sinh ra.\n'
    '3. Kỹ năng AI Engineering: Biết cách prompt chuẩn và tích hợp API mô hình AI vào sản phẩm.\n'
    '4. Tư duy sản phẩm (Product Mindset): Hiểu người dùng cần gì để giải quyết đúng vấn đề.'
)
set_para_font(p5_5, name='Times New Roman', size=12)

# 5.6
h5_6 = doc.add_heading('5.6. Năng lực cần phát triển trong thời đại AI & Vì sao quan trọng? (Slide 53 - Câu 4)', level=2)
set_para_font(h5_6, name='Times New Roman', size=12, bold=True)
p5_6 = doc.add_paragraph(
    '• Tư duy phản biện: Luôn kiểm chứng lại kết quả của AI, không tin tưởng mù quáng.\n'
    '• Kỹ năng bóc tách vấn đề: Chia bài toán lớn thành các phần việc nhỏ, rõ ràng để giao cho AI giải quyết.\n'
    '• Khả năng tự học nhanh: Công nghệ đổi liên tục, kỹ năng quan trọng nhất là học nhanh công cụ mới.\n'
    '• Đạo đức và an toàn thông tin: Có ý thức bảo vệ dữ liệu người dùng và chịu trách nhiệm với sản phẩm của mình.'
)
set_para_font(p5_6, name='Times New Roman', size=12)

# Save document
out_docx = os.path.join(os.path.dirname(__file__), '..', '..', 'TaiLieu', '049206001215_VoVanKhuongThinh_BaiTapTuan1.docx')
doc.save(out_docx)
print(f'Successfully updated docx at: {out_docx}')
